"""v1.4.1+: illegal-char fix, exact session count, prompt rules, LIL yellow-cell guard, notes wiring."""
import io
import json
import re
import sys
import types
from pathlib import Path


class _CM:
    def __enter__(self): return self
    def __exit__(self, *a): return False


class _Proxy:
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def __call__(self, *a, **k): return _CM()
    def __getattr__(self, name): return _Proxy()


def _selector(*args, **kwargs):
    opts = args[1] if len(args) > 1 and isinstance(args[1], (list, tuple)) else None
    if opts:
        idx = kwargs.get("index", 0) or 0
        return opts[idx] if isinstance(idx, int) and idx < len(opts) else opts[0]
    return _Proxy()


fake = types.ModuleType("streamlit")


class _SG(dict):
    def __getattr__(self, n): return _Proxy()
    def __setattr__(self, n, v): dict.__setitem__(self, n, v)
    def setdefault(self, k, d=None):
        if k not in self: dict.__setitem__(self, k, d)
        return dict.__getitem__(self, k)
    def pop(self, k, *a): return dict.pop(self, k, *a)


fake.session_state = _SG()
fake.cache_data = lambda f=None, **k: (f if f else (lambda **kk: None))
fake.secrets = {}
fake.__getattr__ = lambda name: _Proxy()
fake.selectbox = _selector
fake.radio = _selector
fake.tabs = lambda *a, **k: [_CM() for _ in (a[0] if a and isinstance(a[0], (list, tuple)) else [])]
fake.columns = lambda *a, **k: [_CM() for _ in (range(a[0]) if a and isinstance(a[0], int) else (a[0] if a and isinstance(a[0], (list, tuple)) else [1]))]
fake.text_input = lambda *a, **k: ""
fake.text_area = lambda *a, **k: ""
fake.file_uploader = lambda *a, **k: None
fake.button = lambda *a, **k: False
fake.checkbox = lambda *a, **k: False
fake.download_button = lambda *a, **k: False
fake.form_submit_button = lambda *a, **k: False
fake.number_input = lambda *a, **k: 0
fake.slider = lambda *a, **k: 0
sys.modules["streamlit"] = fake
err = types.ModuleType("streamlit.errors")


class _E(Exception):
    pass


err.StreamlitSecretNotFoundError = _E
sys.modules["streamlit.errors"] = err
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import app  # noqa: E402

failures = []


def check(name, condition, extra=""):
    print(f"{'PASS' if condition else 'FAIL'} {name} {extra}")
    if not condition:
        failures.append(name)


# --- 1. The user's exact Mistral error string ---
BAD = ("Integrate Mathematics by having learners calculate energy efficiency of their designs "
       "using formulas like ( \x0cext{Efficiency} = \x0crac{ \x0cext{Useful Energy Output}}"
       "{ \x0cext{Total Energy Input}} \x0cimes 100% ).")
cleaned = app.clean_cell_text(BAD)
check("clean removes form feeds", "\x0c" not in cleaned)
check("clean keeps visible text", "Efficiency" in cleaned and "Useful Energy Output" in cleaned)
check("clean keeps newlines", app.clean_cell_text("line1\nline2\ttab") == "line1\nline2\ttab")
check("clean strips all illegal ranges",
      app.clean_cell_text("".join(chr(i) for i in range(0, 32)) + "ok")
      == "\t\n\rok")  # keeps only Excel-legal control chars: tab, newline, carriage return
check("cell_text sanitizes", "\x0c" not in app.cell_text(BAD))
check("cell_text list joins", "\x0c" not in app.cell_text([BAD, "x"]))
# full legacy path: the string that crashed excel_export before
wb_test = __import__("openpyxl").Workbook()
ws_test = wb_test.active
try:
    ws_test["A1"] = app.cell_text(BAD)
    ok_write = True
except Exception:
    ok_write = False
check("cell_text output is Excel-writable", ok_write)
check("format_flow_text sanitizes", "\x0c" not in app.format_flow_text("Activity: " + BAD))
check("format_references_text sanitizes", "\x0c" not in app.format_references_text(BAD))

# --- 2. Prompt rule updates ---
d = {"area": "Science", "grade": "G9", "term": "Term 1", "week": "Week 3", "strategy": "7Es Model",
     "title": "", "sessions": 5, "duration": "60 minutes", "medium": "English", "teacher": "T. Chua",
     "context": "", "bow": "sample bow text"}
p = app.make_prompt(d)
check("prompt pins exact session count", "exactly 5 learner-centered sessions" in p and "Session 5" in p)
check("prompt forbids fewer sessions", "no fewer, no more" in p)
check("prompt Bloom cap", "Bloom" in p and "beyond the Bloom" in p)
check("prompt KSA", "Knowledge, Skills, and Attitude (KSA)" in p)
check("prompt assessment addresses objectives", "MUST directly address the learning objectives" in p)
check("prompt ways forward higher-level", "HIGHER-LEVEL activity or enhancement" in p)
lp = app.make_lil_prompt({**d, "termweek": "Term 1 · Week 3", "exemplar": "EXEMPLAR"})
check("LIL prompt thorough exemplar", "READ THE ENTIRE LESSON EXEMPLAR THOROUGHLY" in lp)
check("LIL prompt primary-source rule", "PRIMARY SOURCE RULE: Use the Lesson Exemplar as the primary source" in lp
      and "preserve the original activity names, sequence, and content" in lp
      and "not supported by the exemplar" in lp)
check("LIL schema pins exemplar activities", "original names, sequence, and content" in json.dumps(app.LIL_SCHEMA))
check("LIL prompt Bloom/KSA", "Bloom" in lp and "KSA" in lp)
check("LIL prompt assessment alignment", "MUST directly address the session's learning objectives" in lp)
check("LIL prompt higher-level enrichment", "HIGHER-LEVEL activity" in lp)
check("LIL prompt exact sessions", "exactly 5 session object(s)" in lp and "Session 5" in lp)

# --- 3. Session-count enforcement ---
GOOD = {"sessions": [{"session": f"Session {i}", "topic": f"T{i}"} for i in range(1, 6)]}
SHORT = {"sessions": [{"session": "Session 1", "topic": "T1"}]}
OVER = {"sessions": [{"session": f"Session {i}", "topic": f"T{i}"} for i in range(1, 9)]}
orig_ask = app.ask_ai
calls = []
try:
    app.ask_ai = lambda *a, **k: (calls.append(a[0]), json.dumps(GOOD))[1]
    fixed = app._enforce_session_count(json.loads(json.dumps(SHORT)), 5, "retry-prompt")
    check("enforce retries short plan to 5", len(fixed["sessions"]) == 5, f"got {len(fixed['sessions'])}")
    check("retry prompt sent", len(calls) == 1 and calls[0] == "retry-prompt")
    calls.clear()
    same = app._enforce_session_count(json.loads(json.dumps(GOOD)), 5, "retry-prompt")
    check("enforce keeps exact plan, no retry", len(same["sessions"]) == 5 and calls == [])
    calls.clear()
    app.ask_ai = lambda *a, **k: json.dumps(SHORT)  # retry also fails
    padded = app._enforce_session_count(json.loads(json.dumps(SHORT)), 5, "retry-prompt")
    check("enforce pads after failed retry", len(padded["sessions"]) == 5 and padded["sessions"][4]["topic"] == "N/A")
    app.ask_ai = lambda *a, **k: json.dumps(GOOD)
    trimmed = app._enforce_session_count(json.loads(json.dumps(OVER)), 5, "retry-prompt")
    check("enforce truncates over-plan", len(trimmed["sessions"]) == 5)
finally:
    app.ask_ai = orig_ask

# --- 4. LIL export: yellow cells never touched, session picker, bullet ways forward ---
def make_lil_plan():
    return {
        "component": "Curriculum content",
        "learning_competency": ["Dress up according to target work situation"],
        "sessions": [
            {"session": "Session 1", "topic": ["T1a", "T1b", "T1c"],
             "learning_objectives": ["- Obj A1", "- Obj A2", "- Obj A3"],
             "learning_resources": ["- Exemplar\n- Book: Page 12", "- R2", "- R3"],
             "flow": ["Preparation: phase one plan\nPresentation: phase two plan", "F1b", "F1c"],
             "learning_experience": ["- LE1", "- LE2", "- LE3"],
             "assessing_learning": ["- Quiz 1", "- Q2", "- Q3"],
             "ways_forward": ["- Proceed as planned", "- Enrichment", "- Reteach"],
             "worked_well": ["- W1", "- W2", "- W3"],
             "remediation": ["- Rem1", "- Rem2", "- Rem3"],
             "enrichment": ["- En1", "- En2", "- En3"],
             "adjustments": ["- Adj1", "- Adj2", "- Adj3"]},
            {"session": "Session 2", "topic": ["T2a", "T2b", "T2c"],
             "learning_objectives": ["- Obj B1", "- Obj B2", "- Obj B3"],
             "learning_resources": ["- R2x", "- R2y", "- R2z"],
             "flow": ["Preparation: phase one session two\nPresentation: phase two session two", "F2b", "F2c"],
             "learning_experience": ["- LE2x", "- LE2y", "- LE2z"],
             "assessing_learning": ["- Quiz 2", "- Q2b", "- Q2c"],
             "ways_forward": ["- Remediation: follow-up drills", "- Enrichment 2", "- Reteach 2"],
             "worked_well": ["- W2x", "- W2y", "- W2z"],
             "remediation": ["- Rem2x", "- Rem2y", "- Rem2z"],
             "enrichment": ["- En2x", "- En2y", "- En2z"],
             "adjustments": ["- Adj2x", "- Adj2y", "- Adj2z"]},
        ],
    }


DETAILS = {"area": "Science", "grade": "Grade 9", "teacher": "Jose Dennis P. Chua",
           "termweek": "Term 1 · Week 3", "strategy": "5Ps Model", "sessions": 2}

if app.LIL_TEMPLATE.exists():
    from openpyxl import load_workbook
    plan = make_lil_plan()
    blob0 = app.lil_export(plan, DETAILS, {})
    check("LIL export returns bytes", isinstance(blob0, bytes) and len(blob0) > 5000)
    ws0 = load_workbook(io.BytesIO(blob0)).active
    check("exported Term/Week joined", "Term 1 · Week 3" in str(ws0["E11"].value))
    check("teacher on log", "Jose Dennis" in str(ws0["E10"].value) and "Jose Dennis" in str(ws0["A31"].value))
    check("dates/time still blank", ws0["G11"].value in (None, ""))
    text0 = str(ws0["C21"].value)
    check("session 1 flow in column C", "phase one plan" in text0 and "Presentation:" in text0)
    check("session 2 flow in column D (all sessions, one file)", "session two" in str(ws0["D21"].value))
    check("session 2 objectives in column D", "Obj B1" in str(ws0["D18"].value))
    check("session 2 assessment in column D", "Quiz 2" in str(ws0["D22"].value))
    check("day/date labels untouched", str(ws0["C16"].value).strip() == "Day:" and str(ws0["C17"].value).strip() == "Date:")
    check("reminder block preserved", "REMINDER" in str(ws0["F30"].value))
    # yellow label cells keep their original values
    check("yellow A14 untouched", str(ws0["A14"].value) == "Component")
    check("yellow A20 untouched", "WEEKLY IMPLEMENTATION LOG" in str(ws0["A20"].value))
    check("yellow A8 untouched", str(ws0["A8"].value) == "LESSON IMPLEMENTATION LOG")
    check("yellow A23 guide kept", "Proceed as planned" in str(ws0["A23"].value))
    # guard really flags yellow vs blank cells
    try:
        cell_a8 = ws0["A8"]
        fill = cell_a8.fill
        is_yellow = (fill is not None and fill.patternType == "solid"
                     and getattr(fill.start_color, "rgb", None) not in (None, "00000000", "FFFFFFFF"))
        check("yellow detection flags A8", is_yellow, str(getattr(fill.start_color, 'rgb', None)))
        cell_blank = ws0["C21"]
        fb = cell_blank.fill
        not_yellow = not (fb is not None and fb.patternType == "solid"
                          and getattr(fb.start_color, "rgb", None) not in (None, "00000000", "FFFFFFFF"))
        check("blank C21 not flagged yellow", not_yellow)
    except Exception as exc:
        check("yellow detection runs", False, str(exc))
    # ways_forward bullets normalized
    check("ways forward bullet form", "• Proceed as planned" in str(ws0["C23"].value))
    # illegal chars can never reach the file
    plan["sessions"][0]["flow"] = ["Preparation: bad \x0crac here\nPresentation: ok \x0b vert"]
    blob_bad = app.lil_export(plan, DETAILS, {})
    ws_bad = load_workbook(io.BytesIO(blob_bad)).active
    check("illegal chars sanitized in export", "\x0c" not in str(ws_bad["C21"].value) and "\x0b" not in str(ws_bad["C21"].value))
    check("flow still bullets after sanitize", "Preparation:" in str(ws_bad["C21"].value) and "Presentation:" in str(ws_bad["C21"].value))
else:
    check("LIL template present", False)

# --- 5. Additional-instructions (note) wiring on every tab ---
d_note = {**d, "note": "Use farming scenarios and keep sentences short."}
pn = app.make_prompt(d_note)
check("ILAW prompt carries note", "Use farming scenarios" in pn)
check("ILAW prompt none-default", "Extra teacher instructions" in p and "None" in p.split("Extra teacher instructions")[1][:80])
lpn = app.make_lil_prompt({**d, "termweek": "Term 1 · Week 3", "exemplar": "EX", "note": "Base it on the second lesson."})
check("LIL prompt carries note", "Base it on the second lesson." in lpn)
check("LIL prompt none-default", "None" in app.make_lil_prompt({**d, "termweek": "T", "exemplar": "E"}).split("Extra teacher instructions")[1][:80])
tpn = app.make_test_prompt("basis text", {**d, "items": 20, "hots_min": 6, "test_type": "Examination", "mix": {"lots": 40, "mots": 30, "hots": 30}, "note": "Add easy warm-up items."})
check("Test prompt carries note", "Add easy warm-up items." in tpn)
check("Test prompt none-default", "None" in app.make_test_prompt("b", {**d, "items": 20, "hots_min": 6, "test_type": "Quiz", "mix": {"lots": 40, "mots": 30, "hots": 30}}).split("Extra teacher instructions")[1][:80])
ppn = app.make_ppt_prompt("basis", {"slides": 16, "session_number": 1, "session_topic": "T", "note": "Cebuano keywords please."})
check("PPT prompt carries note", "Cebuano keywords please." in ppn)

# --- 6. Version ---
check("version is 1.7.0", app._APP_VERSION == "1.7.0", app._APP_VERSION)

print()
if failures:
    print(f"{len(failures)} FAILED: {failures}")
    sys.exit(1)
print("ALL CHECKS PASSED")
