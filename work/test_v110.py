"""Offline tests for the v1.1.0 changes: options grid, variant picking, doc readers, PPT builder.

Run with the project venv. Streamlit is stubbed so app.py can be imported without a server.
"""
import io
import json
import sys
import types
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # project root so `import app` works

# ---------------------------------------------------------------- streamlit stub
class _CM:
    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def __call__(self, *args, **kwargs):
        return _CM()

    def __getattr__(self, name):
        return _CM()

    def __iter__(self):
        return iter((_CM(), _CM(), _CM()))

    def __bool__(self):
        return False


class _State(dict):
    def __getattr__(self, name):
        try:
            return self[name]
        except KeyError:
            state = _State()
            self[name] = state
            return state

    def __setattr__(self, name, value):
        self[name] = value

    def setdefault(self, key, default=None):
        if key not in self or self[key] is None:
            self[key] = default
        return self[key]


class _SidebarCM(_CM):
    pass


stub = types.ModuleType("streamlit")
stub.set_page_config = lambda *a, **k: None
stub.sidebar = _SidebarCM()
stub.session_state = _State()
stub.secrets = {}  # st.secrets["GOOGLE_API_KEY"] lookups fail gracefully
stub.tabs = lambda labels: tuple(_CM() for _ in labels)
stub.columns = lambda n, **k: tuple(_CM() for _ in range(n if isinstance(n, int) else len(n)))
stub.expander = lambda *a, **k: _CM()
stub.container = lambda **k: _CM()
stub.form = lambda *a, **k: _CM()
stub.spinner = lambda *a, **k: _CM()
stub.title = lambda *a, **k: None
stub.caption = lambda *a, **k: None
stub.markdown = lambda *a, **k: None
stub.info = lambda *a, **k: None
stub.write = lambda *a, **k: None
stub.error = lambda *a, **k: None
stub.success = lambda *a, **k: None
stub.warning = lambda *a, **k: None
stub.button = lambda *a, **k: False
stub.subheader = lambda *a, **k: None
stub.header = lambda *a, **k: None
stub.subheader = lambda *a, **k: None
stub.selectbox = lambda *a, **k: (a[1] or [""])[0] if len(a) > 1 and a[1] else k.get("options", [""])[0]
stub.select_slider = lambda *a, **k: (k.get("options") or [""])[0]
stub.radio = lambda *a, **k: (a[1] or [""])[0]
stub.text_input = lambda *a, **k: (k.get("value") or "") or ""
stub.text_area = lambda *a, **k: (k.get("value") or "") or ""
stub.file_uploader = lambda *a, **k: None
stub.download_button = lambda *a, **k: None
stub.link_button = lambda *a, **k: None
stub.cache_data = lambda f=None, **k: (f if f else (lambda g: g))
stub.form_submit_button = lambda *a, **k: False
stub.rerun = lambda *a, **k: None
stub.link_button = lambda *a, **k: None
stub.cache_data = lambda f=None, **k: (f if f else (lambda g: g))
stub.errors = types.ModuleType("streamlit.errors")
stub.errors.StreamlitSecretNotFoundError = type("StreamlitSecretNotFoundError", (Exception,), {})
sys.modules["streamlit"] = stub
sys.modules["streamlit.errors"] = stub.errors

import app  # noqa: E402

failures = []


def check(name, condition, extra=""):
    print(f"{'PASS' if condition else 'FAIL'} {name} {extra}")
    if not condition:
        failures.append(name)


# ------------------------------------------------------- 1. _ensure_three_options
options = app._ensure_three_options(["only one"], "flow")
check("1. one option padded to three", options == ["only one", "only one", "only one"], options)
options = app._ensure_three_options(["a", "b", "c", "d"], "flow")
check("2. four options trimmed to three", options == ["a", "b", "c"], options)
options = app._ensure_three_options("", "pre_lesson", fallback="F")
check("3. empty cell gets fallback tripled", options == ["F", "F", "F"], options)

# ------------------------------------------------------- 2. first_option helper
check("4. first_option on list", app.first_option(["A", "B", "C"]) == "A")
check("5. first_option on string", app.first_option("plain") == "plain")
check("6. first_option on empty list", app.first_option([], "fb") == "fb")

# ------------------------------------------------------- 3. show_plan grid wiring
app.st.session_state.clear()
plan = {
    "lesson_title": ["Title opt 1", "Title opt 2"],
    "overview": ["Over opt 1", "Over opt 2"],
    "standards_and_competency": ["Std 1"],
    "sessions": [{
        "session": "Session 1",
        "topic": ["Topic 1", "Topic 2", "Topic 3"],
        "learning_objectives": ["LO a", "LO b", "LO c"],
        "pre_lesson": ["P1", "P2", "P3"],
        "flow": ["F1", "F2", "F3"],
        "learning_resources": ["R1", "R2", "R3"],
        "integration": ["I1", "I2", "I3"],
        "formative_assessment": ["FA1", "FA2", "FA3"],
        "extended_learning": ["E1", "E2", "E3"],
        "reflection": ["RE1", "RE2", "RE3"],
    }],
}
details = {"area": "Science", "teacher": "T", "grade": "G9", "week": "W3", "sessions": 1, "term": "Term 1",
           "reference_source": "Budget of Work (BOW) PDF uploaded by teacher.", "ai_provider": "Mistral",
           "context": "", "bow_filename": ""}
try:
    app.show_plan(plan)
    check("7. show_plan grid runs with 3-option cells", True)
except Exception as exc:
    check("7. show_plan grid runs with 3-option cells", False, repr(exc))

# picks default to first option → export uses option 1 content (legacy shape still accepted)
try:
    workbook_bytes = app.excel_export(plan, details, {})
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(workbook_bytes))["WEEKLY LESSON PLAN"]
    check("8. excel export defaults to option 1 (flow)", sheet["B23"].value == "F1", repr(sheet["B23"].value))
    check("9. excel export defaults to option 1 (objectives)", "LO a" in str(sheet["B19"].value), repr(sheet["B19"].value))
    check("10. plan-level fields export as text not list", sheet["B8"].value == "Title opt 1", repr(sheet["B8"].value))
    # user picked option 2 for flow (legacy pick dict still honored)
    workbook_bytes = app.excel_export(plan, details, {(0, "flow"): 1})
    sheet = load_workbook(io.BytesIO(workbook_bytes))["WEEKLY LESSON PLAN"]
    check("11. excel export honors user pick (flow option 2)", sheet["B23"].value == "F2", repr(sheet["B23"].value))
except Exception as exc:
    check("8-11. excel export with picks", False, repr(exc))

# ------------------------------------------------------- 4. test variants + picked_test
app.st.session_state.clear()
test = {
    "test_title": "T", "instructions": "I", "competencies": [{"statement": "S", "days": 3}],
    "items": [
        {"number": 1, "competency": "S", "variants": [
            {"question": "Q1a", "choices": {"A": "1", "B": "2", "C": "3", "D": "4"}, "answer": "A", "rationale": "r", "solo_level": "Relational", "cognitive_level": "Applying"},
            {"question": "Q1b", "choices": {"A": "1", "B": "2", "C": "3", "D": "4"}, "answer": "B", "rationale": "r", "solo_level": "Relational", "cognitive_level": "Applying"},
            {"question": "Q1c", "choices": {"A": "1", "B": "2", "C": "3", "D": "4"}, "answer": "C", "rationale": "r", "solo_level": "Relational", "cognitive_level": "Applying"},
        ]},
    ],
}
try:
    app.show_test(test, None)
    check("12. show_test variant grid runs", True)
except Exception as exc:
    check("12. show_test variant grid runs", False, repr(exc))
chosen = app.picked_test(test)
check("13. picked_test defaults to variant A", chosen["items"][0]["question"] == "Q1a")
app.st.session_state["test_picks"][1] = 2
chosen = app.picked_test(test)
check("14. picked_test honors user pick (C)", chosen["items"][0]["question"] == "Q1c")

# variant normalization: 2 variants → padded to 3; flat item → wrapped
app.st.session_state.clear()
raw = {"items": [{"number": 1, "competency": "S", "variants": [
    {"question": "v1", "choices": {"A": "a", "B": "b", "C": "c", "D": "d"}, "answer": "A", "rationale": "x", "solo_level": "Unistructural", "cognitive_level": "Remembering"},
    {"question": "v2", "choices": {"A": "a", "B": "b", "C": "c", "D": "d"}, "answer": "B", "rationale": "x", "solo_level": "Unistructural", "cognitive_level": "Remembering"},
]}]}
class _FakeFile:
    def __init__(self, text):
        self._text = text

    def getvalue(self):
        return self._text.encode("utf-8")


normalizer_input = json.dumps({"items": [{"number": 1, "variants": [
    {"question": "v1", "choices": {"A": "a", "B": "b", "C": "c", "D": "d"}, "answer": "A", "rationale": "x", "solo_level": "Unistructural", "cognitive_level": "Remembering"},
    {"question": "v2", "choices": {"A": "a", "B": "b", "C": "c", "D": "d"}, "answer": "B", "rationale": "x", "solo_level": "Unistructural", "cognitive_level": "Remembering"},
]}]})
# call the same normalization generate_test applies (replicated inline)
plan_json = app._parse_json_lenient(app._extract_json_text(f"```json\n{normalizer_input}\n```"))
items = plan_json.get("items", [])
for item in items:
    raw_variants = item.get("variants") or []
    variants = [v for v in raw_variants if isinstance(v, dict) and v.get("question")]
    while len(variants) < 3:
        variants.append(dict(variants[-1])) if variants else variants.append({"question": "(No variant returned)"})
    item["variants"] = variants[:3]
check("15. two variants padded to three", len(items[0]["variants"]) == 3, len(items[0]["variants"]))

# ------------------------------------------------------- 5. read_any_document formats
from docx import Document as _Doc

doc = _Doc()
doc.add_paragraph("Learning Competency: describe the phases of the moon")
doc.add_paragraph("Session 1: Phases of the Moon")
buffer = io.BytesIO()
doc.save(buffer)


class _UpFile:
    def __init__(self, name, data):
        self.name = name
        self._data = data

    def getvalue(self):
        return self._data


check("16. docx read", "moon" in app.read_any_document(_UpFile("plan.docx", buffer.getvalue())))

from openpyxl import Workbook as _WB

wb = _WB()
ws = wb.active
ws["A1"] = "FLOW: students model the moon phases"
buf2 = io.BytesIO()
wb.save(buf2)
check("17. xlsx read", "moon" in app.read_any_document(_UpFile("plan.xlsx", buf2.getvalue())))

try:
    app.read_any_document(_UpFile("old.doc", b"x"))
    check("18. doc rejected with clear message", False, "no error raised")
except ValueError as exc:
    check("18. doc rejected with clear message", "docx" in str(exc), str(exc))

# ------------------------------------------------------- 6. PowerPoint builder
if app.Presentation is not None:
    deck_plan = {
        "deck_title": "Phases of the Moon",
        "subject": "Science",
        "theme": {"bg": "F5F9FF", "accent": "2E6FB5", "title": "1A3353", "text": "333333"},
        "slides": [
            {"title": "Phases of the Moon", "bullets": ["Grade 6 Science"], "shape": "none"},
            {"title": "Learning goals", "bullets": ["Identify phases", "Explain cycle"], "shape": "oval"},
            {"title": "New moon", "bullets": ["No light visible"], "shape": "none"},
            {"title": "Full moon", "bullets": ["Fully lit"], "shape": "none"},
        ],
    }
    pptx_bytes = app.build_presentation(deck_plan, "Jose Dennis Plaza Chua")
    check("19. pptx builds", len(pptx_bytes) > 5000, f"{len(pptx_bytes)} bytes")
    check("20. pptx is lightweight (<200 KB)", len(pptx_bytes) < 200_000, f"{len(pptx_bytes)} bytes")
    from pptx import Presentation as _P

    deck = _P(io.BytesIO(pptx_bytes))
    check("21. pptx slide count", len(deck.slides.__iter__.__self__._sldIdLst) == 4 if False else len(deck.slides._sldIdLst) == 4, len(deck.slides._sldIdLst))
    first_texts = "\n".join(shape.text_frame.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    check("22. teacher on title slide", "Jose Dennis Plaza Chua" in first_texts)
else:
    print("SKIP 19-22: python-pptx not installed")

# ------------------------------------------------------- 7. schema sanity
check("23. schema has sessions", "sessions" in json.dumps(app.SCHEMA))
check("24. TEST_SCHEMA is JSON-serializable (no Python sets)", "variants" in json.dumps(app.TEST_SCHEMA))

# ------------------------------------------------------- 8. save_feedback insert-only
import os
import tempfile
import pathlib

tmpdir = tempfile.mkdtemp()

original_feedback = app.FEEDBACK_FILE
app.FEEDBACK_FILE = pathlib.Path(tmpdir) / "feedback.xlsx"
try:
    # 25: local insert-only behaviour (no webhook configured)
    original_hook_file = app.FEEDBACK_HOOK_FILE
    app.FEEDBACK_HOOK_FILE = pathlib.Path(tmpdir) / "none.txt"
    os.environ.pop("FEEDBACK_WEBHOOK_URL", None)
    app.save_feedback("Alice", "⭐⭐⭐⭐⭐", "Great", "Add more")
    app.save_feedback("Bob", "⭐⭐", "Good", "-")
    from openpyxl import load_workbook as _lwb

    ws = _lwb(str(app.FEEDBACK_FILE)).active
    rows = list(ws.iter_rows(values_only=True))
    check("25. feedback header + 2 rows (insert-only)", len(rows) == 3 and rows[1][1] == "Alice" and rows[2][1] == "Bob")

    # 26: hook resolution priority — file line beats env var; comments skipped
    app.FEEDBACK_HOOK_FILE = pathlib.Path(tmpdir) / "hook.txt"
    app.FEEDBACK_HOOK_FILE.write_text(
        "# comment line\n\nhttps://script.google.com/macros/s/FILE_ID/exec\n", encoding="utf-8")
    os.environ["FEEDBACK_WEBHOOK_URL"] = "https://script.google.com/macros/s/ENV_ID/exec"
    check("26a. hook from feedback_webhook.txt wins", app.feedback_hook_url().endswith("FILE_ID/exec"), app.feedback_hook_url())
    app.FEEDBACK_HOOK_FILE = pathlib.Path(tmpdir) / "none.txt"
    check("26b. env var used when file absent", app.feedback_hook_url().endswith("ENV_ID/exec"), app.feedback_hook_url())
    os.environ.pop("FEEDBACK_WEBHOOK_URL")
    check("26c. empty when nothing configured", app.feedback_hook_url() == "")

    # 27: save_feedback returns synced/local-only/sync-failed
    class _Resp:
        def read(self):
            return b"OK"

    class _BadHook:
        pass

    app.FEEDBACK_HOOK_FILE = pathlib.Path(tmpdir) / "hook2.txt"
    app.FEEDBACK_HOOK_FILE.write_text("https://script.google.com/macros/s/X/exec\n", encoding="utf-8")
    original_urlopen = app.urllib.request.urlopen
    app.urllib.request.urlopen = lambda *a, **k: (_ for _ in ()).throw(OSError("no network"))
    status = app.save_feedback("Cara", "⭐⭐⭐", "fine", "-")
    check("27a. network failure -> local save + sync-failed", status.startswith("sync-failed") and len(list(_lwb(str(app.FEEDBACK_FILE)).active.iter_rows(values_only=True))) == 4, status)
    app.urllib.request.urlopen = lambda *a, **k: _Resp()
    status = app.save_feedback("Dan", "⭐⭐⭐⭐", "good", "-")
    check("27b. webhook reachable -> synced", status == "synced", status)
    app.urllib.request.urlopen = original_urlopen
    app.FEEDBACK_HOOK_FILE = pathlib.Path(tmpdir) / "none.txt"

    # 27c (new): no hook -> posts to the owner's Google Form; captured payload check
    captured = {}

    def _fake_urlopen(request, timeout=0):
        captured["url"] = request.full_url
        captured["body"] = request.data.decode("utf-8")
        return _Resp()

    app.urllib.request.urlopen = _fake_urlopen
    status = app.save_feedback("Eve", "⭐⭐", "Form works", "Add charts")
    from urllib.parse import parse_qs as _pqs

    fields = {k: v[0] for k, v in _pqs(captured.get("body", "")).items()}
    check("27c. no hook -> Google Form post succeeds", status == "synced" and captured.get("url", "").endswith("/formResponse"), status)
    check("27d. form payload has correct entries", fields.get(app._FEEDBACK_FORM_ENTRIES["name"]) == "Eve" and fields.get(app._FEEDBACK_FORM_ENTRIES["rating"]) == "2" and fields.get(app._FEEDBACK_FORM_ENTRIES["feedback"]) == "Form works" and fields.get(app._FEEDBACK_FORM_ENTRIES["suggestions"]) == "Add charts", fields)

    # 27e: form unreachable -> local row still saved, honest sync-failed
    app.urllib.request.urlopen = lambda *a, **k: (_ for _ in ()).throw(OSError("offline"))
    status = app.save_feedback("Faye", "⭐⭐⭐", "offline test", "-")
    rows_now = len(list(_lwb(str(app.FEEDBACK_FILE)).active.iter_rows(values_only=True)))
    check("27e. form offline -> local save + sync-failed", status.startswith("sync-failed") and rows_now == 7, f"{status} rows={rows_now}")
    app.urllib.request.urlopen = original_urlopen
finally:
    app.FEEDBACK_FILE = original_feedback

print()
print("ALL V1.1.0 TESTS PASSED" if not failures else f"FAILURES: {failures}")
sys.exit(1 if failures else 0)
