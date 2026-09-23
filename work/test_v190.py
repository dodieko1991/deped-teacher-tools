"""v1.9.0 — Test Paper upload-first: 2–5 ILAW/LIL files, combined basis, auto-detect."""
import sys, types, io
from pathlib import Path

class _CM:
    def __enter__(self): return self
    def __exit__(self, *a): return False

class _Proxy:
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def __call__(self, *a, **k): return _CM()
    def __getattr__(self, name): return _Proxy()

def _streamlit_stub():
    st = types.ModuleType("streamlit")
    errs = types.ModuleType("streamlit.errors")
    class E(Exception): pass
    errs.StreamlitInvalidLayoutContextError = E
    st.errors = errs
    st.__path__ = []
    def _mk(*a, **k): return _CM()
    st.form = _mk; st.form_submit_button = lambda *a, **k: True
    st.columns = lambda n, *a, **k: [_CM() for _ in (range(n) if isinstance(n, int) else n)]
    st.tabs = lambda labels: [_CM() for _ in labels]
    st.sidebar = _Proxy()
    st.session_state = {}
    for name in ("subheader", "info", "warning", "error", "success", "caption", "write",
                 "markdown", "title", "header", "expander", "spinner", "divider", "help",
                 "set_page_config", "image", "progress", "stop", "rerun", "select_slider",
                 "metric", "plotly_chart", "dataframe", "table", "json",
                 "code", "latex", "button", "checkbox", "download_button", "link_button"):
        setattr(st, name, _mk)
    st.cache_data = lambda f=None, **k: (f if f else (lambda g: g))
    st.radio = _mk
    st.file_uploader = lambda *a, **k: []  # multi-file uploader
    st.multiselect = lambda label, options=[], *a, **k: list(options)
    st.text_input = lambda label, value="", *a, **k: ("" if value is None else value)
    st.text_area = lambda label, value="", *a, **k: ("" if value is None else value)
    st.number_input = lambda label, value=0, *a, **k: value
    st.selectbox = lambda label, options=[], *a, **k: (options[0] if options else None)
    st.slider = lambda label, *a, value=None, **k: (value if value is not None else 1)
    st.secrets = {}
    return st

sys.modules["streamlit"] = _streamlit_stub()
if "streamlit.errors" not in sys.modules:
    _errs_mod = types.ModuleType("streamlit.errors")
    class StreamlitSecretNotFoundError(Exception): pass
    _errs_mod.StreamlitSecretNotFoundError = StreamlitSecretNotFoundError
    sys.modules["streamlit.errors"] = _errs_mod

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

failures = []
def check(name, cond, extra=""):
    print(f"{'PASS' if cond else 'FAIL'} {name} {extra}")
    if not cond:
        failures.append(name)

# ------------------------------------------------- ILAW Excel fallback detection
ilaw_xlsx = app.excel_export(
    {"lesson_title": "Newton's Laws", "standards_and_competency": "Force and motion",
     "sessions": [{"session": "SESSION 1", "topic": "t", "learning_objectives": "o", "pre_lesson": "p",
                   "flow": "f", "learning_resources": "- Book, A, Page 1", "integration": "i",
                   "formative_assessment": "fa", "extended_learning": "e", "reflection": "r"}]},
    {"area": "Science", "grade": "Grade 9 – Hydrogen", "teacher": "T", "term": "Term 1", "week": "Week 3",
     "sessions": 1, "strategy": "4Es", "medium": "English", "title": "t", "context": "", "note": "",
     "reference_source": "x", "ai_provider": "Gemini", "bow_filename": "b.pdf", "online_sources": []})
ws = load_workbook(io.BytesIO(ilaw_xlsx))["WEEKLY LESSON PLAN"]
cells = " ".join(str(ws.cell(row=r, column=c).value) for r in range(1, 36) for c in range(1, 7))
meta = app.detect_exemplar_meta(cells)
check("ilaw-xlsx: area", meta["area"].strip().lower() == "science", meta["area"])
check("ilaw-xlsx: grade", meta["grade"].startswith("Grade 9"), meta["grade"])
check("ilaw-xlsx: term/week from Term/Week cell", meta["term"] == "Term 1" and meta["week"] == "Week 3", str(meta))

# junk values never leak in
junk = app.detect_exemplar_meta("Learning Area/s Science None None None Grade Level and Section Grade 9 None")
check("ilaw-xlsx: no 'None' leaks", "None" not in junk["area"] and "None" not in junk["grade"], str(junk))

# real exemplar detection still intact
le_txt = Path(__file__).with_name("le_gm_q1.txt").read_text(encoding="utf-8")
meta2 = app.detect_exemplar_meta(le_txt)
check("le: still detects GM", meta2["area"].lower().replace("&", "and") == "general mathematics", meta2["area"])
check("le: still Grade 11 / Term 1", meta2["grade"] == "Grade 11" and meta2["term"] == "Term 1", str(meta2))

# ------------------------------------------------- Test tab wiring
src = Path(__file__).resolve().parent.parent.joinpath("app.py").read_text(encoding="utf-8")
test_start = src.index("with test_tab:")
test_src = src[test_start: src.index("if test := st.session_state")]
check("order: file uploader before test form", test_src.find("t_files") != -1
      and test_src.find('st.form("test_form")') > test_src.find("t_files"))
check("wiring: accept_multiple_files", "accept_multiple_files=True" in test_src)
check("wiring: 2-5 limit enforced", "2 <= len(t_files) <= 5" in test_src)
check("wiring: files combined into basis", '"\\n\\n".join(t_file_texts)' in test_src)
check("wiring: per-file headers in basis", "=== FILE:" in test_src)
check("wiring: detection per file", test_src.count("detect_exemplar_meta") >= 1)
check("wiring: locks depend on basis choice", "basis_choice == BASIS_FILES" in test_src)
check("wiring: requires >=2 readable files", 'len(t_file_texts) < 2' in test_src)
check("wiring: BASIS_FILES constant", 'BASIS_FILES = "Uploaded ILAW / LIL files (2–5)"' in src)
check("wiring: legacy BASIS_ILAW kept", "BASIS_ILAW" in src.split("BASIS_FILES")[0])
check("wiring: basis_files in details", '"basis_files"' in test_src)
check("order: detection before form", test_src.find("detect_exemplar_meta") < test_src.find('st.form("test_form")'))

# ------------------------------------------------- version
check("version: 1.9.0", app._APP_VERSION == "1.9.0", app._APP_VERSION)

print()
if failures:
    print(f"FAILED {len(failures)}: " + ", ".join(failures))
else:
    print("ALL CHECKS PASSED")
sys.exit(1 if failures else 0)
